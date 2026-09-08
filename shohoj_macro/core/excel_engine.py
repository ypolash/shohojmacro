import os
import openpyxl
from typing import Dict, List, Any, Optional

class ExcelDataEngine:
    """Reads .xlsx files and supports writing status updates."""
    
    def __init__(self):
        self.workbook: Optional[openpyxl.Workbook] = None
        self.sheet = None
        self.filepath = ""
        self.headers: List[str] = []
        
    def load_file(self, filepath: str, sheet_name: str = None) -> bool:
        """Loads an Excel file and initializes headers."""
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"Excel file not found: {filepath}")
            
        self.filepath = filepath
        self.workbook = openpyxl.load_workbook(filepath)
        
        if sheet_name:
            if sheet_name in self.workbook.sheetnames:
                self.sheet = self.workbook[sheet_name]
            else:
                raise ValueError(f"Sheet '{sheet_name}' not found.")
        else:
            self.sheet = self.workbook.active
            
        # Extract headers from the first row
        self.headers = [str(cell.value) if cell.value else f"Col{i+1}" 
                       for i, cell in enumerate(self.sheet[1])]
        return True
        
    def get_row_count(self) -> int:
        """Returns the total number of rows (including header)."""
        if not self.sheet:
            return 0
        return self.sheet.max_row
        
    def get_row_data(self, row_idx: int) -> Dict[str, str]:
        """
        Returns data for a specific row index (1-based, where 1 is header).
        Example: get_row_data(2) gets the first data row.
        """
        if not self.sheet or row_idx < 1 or row_idx > self.sheet.max_row:
            return {}
            
        row_values = [cell.value for cell in self.sheet[row_idx]]
        data = {}
        
        for i, header in enumerate(self.headers):
            if i < len(row_values):
                val = row_values[i]
                data[header] = str(val) if val is not None else ""
            else:
                data[header] = ""
                
        return data
        
    def mark_row_status(self, row_idx: int, status_col_letter: str, status: str):
        """Writes a status message to a specific column on a specific row."""
        if not self.sheet:
            return
            
        cell_ref = f"{status_col_letter}{row_idx}"
        self.sheet[cell_ref] = status
        self.save()
        
    def save(self):
        """Saves changes back to the Excel file."""
        if self.workbook and self.filepath:
            self.workbook.save(self.filepath)
